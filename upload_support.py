"""Candidate SESSION_FILE transport; no model inference and no health-data logging."""
import base64
import csv
import hashlib
import io
import json
import math
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile

FIELDS = (
    "age", "gender", "temperature", "pulse", "systolic", "diastolic",
    "left_naked_eye", "right_naked_eye", "heart_rate", "hear", "exercise", "heart_rhythm",
    "smoke", "HGB", "PLT", "Creatinine", "Urea", "TC", "LDL_C", "HDL_C",
    "hypertension_final", "diabetes_final", "dyslipidemia_final",
    "height_cm", "waist_cm", "weight_kg", "glucose_mmolL", "triglycerides_mmolL",
)
BINARY = {"gender", "hear", "exercise", "heart_rhythm", "smoke", "hypertension_final", "diabetes_final", "dyslipidemia_final"}
POSITIVE = {"height_cm", "waist_cm", "weight_kg", "glucose_mmolL", "triglycerides_mmolL"}
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000


class InputError(ValueError):
    pass


def sheet_names(raw):
    _check_xlsx(raw)
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def _check_xlsx(raw):
    if not raw or len(raw) > MAX_BYTES:
        raise InputError("文件为空或超过5MiB。")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024:
                raise InputError("Excel解压后超过32MiB，请另存为CSV。")
    except zipfile.BadZipFile:
        raise InputError("这不是有效的xlsx文件，请使用Excel另存为xlsx。") from None


def prepare_table(raw, suffix, sheet=None):
    """Validate numbers and send only the 28 model fields, preserving row order."""
    if not raw or len(raw) > MAX_BYTES:
        raise InputError("文件为空或超过5MiB。")
    book = None
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        else:
            raise InputError("CSV编码无法识别，请另存为UTF-8 CSV。")
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    elif suffix == ".xlsx":
        _check_xlsx(raw)
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        if sheet not in book.sheetnames:
            book.close()
            raise InputError("请选择要评估的Excel工作表。")
        reader = book[sheet].iter_rows(values_only=True)
    else:
        raise InputError("仅支持csv和xlsx；旧xls请先另存为xlsx。")
    try:
        header = next(reader, None)
        if header is None:
            raise InputError("表格为空。")
        names = [str(value).strip() if value is not None else "" for value in header]
        if any(names.count(field) > 1 for field in FIELDS):
            raise InputError("模型字段列名重复，请检查表头。")
        missing = [field for field in FIELDS if field not in names]
        if missing:
            raise InputError("缺少列：" + "、".join(missing))
        indices = [names.index(field) for field in FIELDS]
        out = io.StringIO(newline="")
        writer = csv.writer(out, lineterminator="\n")
        writer.writerow(FIELDS)
        count = 0
        for row in reader:
            count += 1
            if count > MAX_ROWS:
                raise InputError("最多支持10000条，请拆分文件。")
            values = []
            for field, index in zip(FIELDS, indices):
                cell = row[index] if index < len(row) else None
                try:
                    if isinstance(cell, bool) or cell is None or str(cell).strip() == "":
                        raise ValueError
                    value = float(cell)
                    if not math.isfinite(value) or (field in BINARY and value not in (0, 1)) or (field in POSITIVE and value <= 0):
                        raise ValueError
                    if field == "age" and value < 65:
                        raise ValueError
                except (ValueError, TypeError, OverflowError):
                    raise InputError(f"第{count}条数据的{field}缺失或不符合模型要求，请检查原表。") from None
                values.append(format(value, ".17g"))
            writer.writerow(values)
        if count == 0:
            raise InputError("表格只有表头，没有数据。")
        payload = out.getvalue().encode("utf-8")
        if len(payload) > MAX_BYTES:
            raise InputError("转换后的CSV超过5MiB，请拆分文件。")
        return payload, count
    except csv.Error:
        raise InputError("CSV结构不正确，请检查分隔符和引号。") from None
    finally:
        if book is not None:
            book.close()


def safe_tag(value):
    value = str(value or "unknown")
    return value if re.fullmatch(r"[A-Za-z0-9_.:-]{1,128}", value) else "unknown"


class UploadError(RuntimeError):
    def __init__(self, stage, status=None, code=None, request_id=None):
        self.stage = stage
        self.status = status
        self.code = safe_tag(code)
        self.request_id = safe_tag(request_id)
        super().__init__(f"{stage}未完成；HTTP={status or '-'}；code={self.code}；request_id={self.request_id}")


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


class FileClient:
    def __init__(self, api_key, workspace, opener=None):
        if not re.fullmatch(r"[A-Za-z0-9-]+", workspace or ""):
            raise InputError("请在Secrets填写有效的BAILIAN_UPLOAD_WORKSPACE_ID。")
        self.base = f"https://{workspace}.cn-beijing.maas.aliyuncs.com/api/v1/connector/dash/"
        self.api_key = api_key
        self.opener = opener or urllib.request.build_opener(NoRedirect())

    def _json(self, operation, body):
        stage = {"applyFileUploadLease": "申请上传地址", "addFile": "注册会话文件", "describeFile": "检查文件状态"}[operation]
        request = urllib.request.Request(self.base + operation, method="POST",
            data=json.dumps(body).encode(), headers={"Authorization": "Bearer " + self.api_key, "Content-Type": "application/json"})
        try:
            with self.opener.open(request, timeout=60) as response:
                result = json.load(response)
        except urllib.error.HTTPError as exc:
            try:
                error = json.loads(exc.read(65536))
            except (ValueError, OSError):
                error = {}
            if not isinstance(error, dict):
                error = {}
            raise UploadError(stage, exc.code, error.get("code"), error.get("requestId") or error.get("request_id")) from None
        except (OSError, ValueError):
            raise UploadError(stage, code="NetworkOrResponseError") from None
        if not isinstance(result, dict):
            raise UploadError(stage, code="UnexpectedResponse")
        if result.get("code") != "Success" or str(result.get("status", result.get("status_code", 200))) != "200":
            raise UploadError(stage, code=result.get("code"), request_id=result.get("requestId"))
        if not isinstance(result.get("data"), dict):
            raise UploadError(stage, code="MissingData")
        return result["data"]

    def upload(self, data):
        # CSV is plain text. A .txt transport name uses a documented upload format;
        # the host must feed these exact CSV bytes to Skill with --format csv.
        lease = self._json("applyFileUploadLease", {"category": "default", "categoryType": "SESSION_FILE",
            "fileName": "cmm_" + uuid.uuid4().hex + ".txt", "sizeBytes": str(len(data)),
            "contentMd5": base64.b64encode(hashlib.md5(data).digest()).decode()})
        param = lease.get("param", {})
        url = param.get("url", "")
        parsed = urllib.parse.urlsplit(url)
        host = parsed.hostname or ""
        # Never forward the API key to the presigned upload destination.
        if parsed.scheme != "https" or not host.endswith(".aliyuncs.com") or parsed.username or parsed.password or parsed.port not in (None, 443) or param.get("method") != "PUT":
            raise UploadError("检查上传地址", code="UnexpectedUploadDestination")
        upload_headers = dict(param.get("headers", {}))
        if any(name.lower() in {"authorization", "proxy-authorization", "host"} for name in upload_headers):
            raise UploadError("检查上传地址", code="UnexpectedUploadHeaders")
        request = urllib.request.Request(url, data=data, method="PUT", headers=upload_headers)
        try:
            with self.opener.open(request, timeout=120) as response:
                if response.status not in (200, 201, 204):
                    raise UploadError("传输文件", response.status)
        except urllib.error.HTTPError as exc:
            raise UploadError("传输文件", exc.code) from None
        except OSError:
            raise UploadError("传输文件", code="NetworkError") from None
        lease_id = lease.get("leaseId") or lease.get("fileUploadLeaseId")
        if not lease_id:
            raise UploadError("注册会话文件", code="MissingLeaseId")
        added = self._json("addFile", {"leaseId": lease_id, "category": "default", "categoryType": "SESSION_FILE", "parser": "AUTO_SELECT"})
        file_id = added.get("fileId", "")
        if not re.fullmatch(r"file_session_[A-Za-z0-9_-]+", file_id):
            raise UploadError("注册会话文件", code="ExpectedSessionFileId")
        return file_id

    def status(self, file_id):
        result = self._json("describeFile", {"fileId": file_id})
        status = result.get("status")
        if status in ("PARSE_SUCCESS", "FILE_IS_READY"):
            return True
        if status in ("FAIL", "PARSE_FAILED", "FILE_PARSE_FAILED"):
            raise UploadError("解析会话文件", code="FileParseFailed")
        return False


def conversation_body(prompt, session_id=None, file_id=None):
    input_data = {"prompt": prompt}
    if session_id:
        input_data["session_id"] = session_id
    parameters = {}
    if file_id:
        parameters["session_file_ids"] = [file_id]
        input_data["prompt"] = prompt + (
            "\n\n本次附有用户授权的体检表，传输格式为UTF-8 CSV纯文本（文件名后缀为.txt）。"
            "请由宿主定位本次附件的原始沙箱文件，按CMM Skill v1.2的init/stage/run流程计算，stage使用--format csv。"
            "不要读取或打印整表，不把原始数据或逐人结果发送到任何MCP。"
            "如果当前API未把附件提供为可用的原始沙箱文件，请明确回复‘未获得原始附件，尚未执行CMM计算’，"
            "不得根据解析文本重建表格、心算概率或声称计算完成。仅真实完成后报告处理条数和文件交付情况。"
        )
    return {"input": input_data, "parameters": parameters}
